import io, os, re, sys, zipfile
from pptx import Presentation
from PIL import Image
SRC=sys.argv[1] if len(sys.argv) > 1 else sys.exit("usage: make_team_template.py <team-deck.pptx> [out.pptx]")
OUT=sys.argv[2] if len(sys.argv) > 2 else os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "assets", "team-template.pptx")
prs=Presentation(SRC)
sldIdLst=prs.slides._sldIdLst
for sldId in list(sldIdLst)[1:]:
    prs.part.drop_rel(sldId.rId); sldIdLst.remove(sldId)
s=prs.slides[0]
by={}
for sh in list(s.shapes):
    if sh.has_text_frame and sh.text_frame.text.strip()=="XXX":
        sh._element.getparent().remove(sh._element); continue
    by[sh.name]=sh
def set_para(p,text):
    runs=p.runs
    runs[0].text=text
    for r in runs[1:]: r._r.getparent().remove(r._r)
tf=by["Title 3"].text_frame
set_para(tf.paragraphs[0],"Presentation title"); set_para(tf.paragraphs[1],"One-line subtitle")
tf=by["Text Placeholder 4"].text_frame
set_para(tf.paragraphs[0],"Presenter Name")
for p in tf.paragraphs[1:]: p._p.getparent().remove(p._p)
set_para(by["Text Placeholder 5"].text_frame.paragraphs[0],"Team Name")
if s.has_notes_slide: s.notes_slide.notes_text_frame.text=""
cp=prs.core_properties
for k in ("author","last_modified_by","title","subject","comments","keywords","category"):
    setattr(cp,k,"")
buf=io.BytesIO(); prs.save(buf)
# replace thumbnail with a blank image
thumb=io.BytesIO(); Image.new("RGB",(256,144),"white").save(thumb,"JPEG")
zin=zipfile.ZipFile(io.BytesIO(buf.getvalue()))
with zipfile.ZipFile(OUT,"w",zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data=zin.read(item.filename)
        if item.filename=="docProps/thumbnail.jpeg": data=thumb.getvalue()
        if item.filename=="docProps/app.xml":

            x=data.decode("utf8")
            x=re.sub(r"<Slides>\d+</Slides>","<Slides>1</Slides>",x)
            x=re.sub(r"<Notes>\d+</Notes>","<Notes>0</Notes>",x)
            x=re.sub(r"(<vt:lpstr>Slide Titles</vt:lpstr></vt:variant><vt:variant><vt:i4>)\d+","\\g<1>1",x)
            parts=re.search(r'<TitlesOfParts><vt:vector size="(\d+)"',x)
            n=int(parts.group(1)); extra=x.count("<vt:lpstr>PowerPoint Presentation</vt:lpstr>")-1
            x=x.replace("<vt:lpstr>PowerPoint Presentation</vt:lpstr>","",extra)
            x=x.replace(f'<TitlesOfParts><vt:vector size="{n}"',f'<TitlesOfParts><vt:vector size="{n-extra}"')
            x=re.sub(r"<Template>[^<]*</Template>","<Template></Template>",x)
            x=re.sub(r"<TotalTime>\d+</TotalTime>","<TotalTime>0</TotalTime>",x)
            data=x.encode("utf8")
        zout.writestr(item,data)
print("ok")
