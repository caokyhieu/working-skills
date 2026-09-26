#!/usr/bin/env python3
"""Build editable decks in the team theme (see references/team-theme.md).

The base template (assets/team-template.pptx) supplies the masters, branded header,
logo, confidentiality line, page number, and title-slide artwork. This module adds
native, editable shapes that follow the team's component vocabulary.

    from team_deck import TeamDeck
    deck = TeamDeck(kicker="Research Insight: Learned Cardinality Estimation")
    deck.title_slide("Title", "Subtitle", presenters=["Name"], team="Intelligent Data Team")
    s = deck.slide("Conclusion-led title", tagline="optional one-line context")
    s.section(0.78, 1.15, 5.5, "HOW IT WORKS")
    s.card(0.78, 1.45, 5.5, 1.8, "1 · Baseline", ["point one", "point two"])
    s.takeaway("Bold lead sentence.", "Supporting detail.")
    s.notes("Speaker notes with sources.")
    deck.save("out/deck-v1.pptx")        # prints layout warnings

Coordinates are inches on a 13.333 x 7.5 in slide. Requires python-pptx.
"""

import copy
import math
import os
import sys

from lxml import etree
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "..", "assets", "team-template.pptx")

# ---------------------------------------------------------------- tokens
FONT = "Microsoft YaHei"
C = {
    "title": "C00000",      # slide titles, stat numbers, emphasis text
    "red": "C7000B",        # fills, accent borders, focus series, arrows that matter
    "kicker": "0070C0",     # deck/topic kicker above the title
    "corner": "00B0F0",     # top-left corner triangle
    "ink": "1D1D1A",        # headings, dark panels
    "black": "000000",      # strong body text on light fills
    "body": "595757",       # explanatory text
    "muted": "898989",      # section labels, captions, neutral arrows
    "surface": "F5F5F5",    # default panel fill
    "white": "FFFFFF",
    "border": "DDDDDD",     # thin panel borders, rules
    "track": "F0F0F0",      # bar tracks
    "grey_bar": "A5A5A3",   # secondary/baseline bars
    "pale_red": "FDF0EF",   # soft accent fill
    "pink": "E5B4B0",       # light accent bars
    "on_dark_sub": "C9CEDA",
    "on_red_sub": "F3D3D2",
}
SLIDE_W, SLIDE_H = 13.333, 7.5
MARGIN_X, CONTENT_W = 0.78, 11.78
CONTENT_TOP, CONTENT_TOP_TAGLINE, CONTENT_BOTTOM = 1.10, 1.30, 6.85
LOGO_ZONE = (10.7, 6.8)          # x >= and y >= : master logo area, keep clear
MIN_PT = 7.5

# card variants: fill, border, border width pt, dash, title color, body color
VARIANTS = {
    "neutral": ("surface", "border", 0.75, False, "ink", "body"),
    "white":   ("white", "border", 0.75, False, "ink", "body"),
    "dark":    ("ink", "ink", 0.75, False, "white", "on_dark_sub"),
    "accent":  ("red", None, 0, False, "white", "on_red_sub"),
    "outline": ("white", "red", 1.0, False, "title", "body"),     # emphasized / key step
    "dashed":  ("white", "red", 1.0, True, "title", "body"),      # inferred, future, in progress
    "pale":    ("pale_red", "red", 0.75, False, "title", "body"),
    "plain":   (None, None, 0, False, "ink", "body"),
}


def rgb(key_or_hex):
    return RGBColor.from_string(C.get(key_or_hex, key_or_hex))


def _font(run, size, color, bold=False, spacing=None, italic=False):
    f = run.font
    f.size = Pt(size)
    f.bold = bold
    f.italic = italic
    f.name = FONT
    f.color.rgb = rgb(color)
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", FONT)
    if spacing is not None:
        rPr.set("spc", str(int(spacing)))


def _line_spacing(paragraph, pct):
    pPr = paragraph._p.get_or_add_pPr()
    for old in pPr.findall(qn("a:lnSpc")):
        pPr.remove(old)
    ln = etree.SubElement(pPr, qn("a:lnSpc"))
    etree.SubElement(ln, qn("a:spcPct")).set("val", str(int(pct * 1000)))
    pPr.insert(0, ln)


def _frame(tf, anchor="t", wrap=True):
    tf.auto_size = MSO_AUTO_SIZE.NONE  # keep boxes fixed; autofit re-centres single-line text
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "ctr": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]


def _write(tf, paragraphs, size, color, bold=False, align="l", line=1.3, spacing=None):
    """paragraphs: str | list of (str | list of (text, {overrides}))."""
    if isinstance(paragraphs, str):
        paragraphs = [paragraphs]
    al = {"l": PP_ALIGN.LEFT, "ctr": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
    for i, para in enumerate(paragraphs):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = al
        _line_spacing(p, line * 100)
        runs = [(para, {})] if isinstance(para, str) else para
        for text, o in runs:
            r = p.add_run()
            r.text = text
            _font(r, o.get("size", size), o.get("color", color), o.get("bold", bold), o.get("spacing", spacing),
                  o.get("italic", False))


class Canvas:
    """One content slide. All methods return the created shape."""

    def __init__(self, deck, slide, name):
        self.deck, self.slide, self.name = deck, slide, name
        self._text_checks = []

    # -- primitives ----------------------------------------------------
    def rect(self, x, y, w, h, fill="surface", border="border", border_pt=0.75, dashed=False, shape=MSO_SHAPE.RECTANGLE):
        sh = self.slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
        if fill:
            sh.fill.solid()
            sh.fill.fore_color.rgb = rgb(fill)
        else:
            sh.fill.background()
        if border:
            sh.line.color.rgb = rgb(border)
            sh.line.width = Pt(border_pt)
            if dashed:
                ln = sh.line._get_or_add_ln()
                etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
        else:
            sh.line.fill.background()
        style = sh._element.find(qn("p:style"))
        if style is not None:  # theme style adds shadows/effects the team decks don't use
            sh._element.remove(style)
        if sh.has_text_frame:
            sh.text_frame.text = ""
        return sh

    def text(self, x, y, w, h, paragraphs, size=12, color="body", bold=False, align="l", anchor="t", line=1.3, spacing=None, wrap=True):
        tb = self.slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        _frame(tb.text_frame, anchor, wrap)
        _write(tb.text_frame, paragraphs, size, color, bold, align, line, spacing)
        self._text_checks.append((tb, x, y, w, h, line, wrap))
        return tb

    # -- team components -----------------------------------------------
    def section(self, x, y, w, label, size=10.5):
        """Grey letter-spaced caps label above a region, e.g. 'HOW IT WORKS'."""
        return self.text(x, y, w, 0.22, label.upper(), size=size, color="muted", bold=True, anchor="ctr", spacing=200, wrap=False)

    def heading(self, x, y, w, text, size=11, color="ink"):
        return self.text(x, y, w, 0.26, text, size=size, color=color, bold=True, anchor="ctr")

    def card(self, x, y, w, h, title=None, body=None, variant="neutral", tag=None, tag_variant=None, size=9.5, title_size=11, pad=0.2):
        """Panel with optional small tag (e.g. 'NEW', 'BASELINE'), bold title and body lines."""
        fill, border, bpt, dashed, tcol, bcol = VARIANTS[variant]
        shape = self.rect(x, y, w, h, fill, border, bpt, dashed) if (fill or border) else None
        cy = y + pad * 0.7
        if tag:
            self.tag(x + pad, cy, tag, tag_variant or ("accent" if variant == "dark" else "outline"))
            cy += 0.31
        if title:
            self.text(x + pad, cy, w - 2 * pad, 0.26, title, size=title_size, color=tcol, bold=True, anchor="ctr")
            cy += 0.3
        if body:
            lines = [body] if isinstance(body, str) else body
            lines = [("· " + l) if isinstance(l, str) and len(lines) > 1 else l for l in lines]
            self.text(x + pad, cy, w - 2 * pad, max(0.2, y + h - cy - pad * 0.6), lines, size=size, color=bcol, line=1.4)
        return shape

    def node(self, x, y, w, h, label, sub=None, variant="neutral", size=9.5, sub_size=8.0):
        """Centered diagram box: bold label plus optional smaller second line."""
        fill, border, bpt, dashed, tcol, bcol = VARIANTS[variant]
        shape = self.rect(x, y, w, h, fill, border, bpt, dashed)
        if variant in ("neutral", "white"):
            tcol = "black"
        paras = [[(label, {"bold": True, "color": tcol, "size": size})]]
        if sub:
            paras.append([(sub, {"bold": False, "color": bcol, "size": sub_size})])
        self.text(x + 0.06, y, w - 0.12, h, paras, size=size, color=tcol, align="ctr", anchor="ctr", line=1.2)
        return shape

    def tag(self, x, y, label, variant="outline", size=7.5):
        w = max(0.5, 0.075 * len(label) + 0.25)
        if variant == "accent":
            self.rect(x, y, w, 0.21, "red", "red", 0.75)
            col = "white"
        else:
            self.rect(x, y, w, 0.21, "white", "red", 0.75)
            col = "title"
        self.text(x, y, w, 0.21, label.upper(), size=size, color=col, bold=True, align="ctr", anchor="ctr", wrap=False)
        return w

    def strip(self, x, y, w, h, lead, rest="", size=9.5):
        """Grey callout strip with a red left bar: bold lead + grey detail."""
        self.rect(x, y, w, h, "surface", "border", 0.75)
        self.rect(x, y, 0.04, h, "red", None)
        runs = [(lead, {"bold": True, "color": "ink"})]
        if rest:
            runs.append((" " + rest, {"color": "body"}))
        return self.text(x + 0.25, y, w - 0.4, h, [runs], size=size, anchor="ctr", line=1.4)

    def takeaway(self, lead, rest="", y=None, h=0.45, size=12, rule=True):
        """Bottom conclusion line (bold lead + detail), optionally with a rule above."""
        y = y if y is not None else 6.3
        if rule:
            self.rule(MARGIN_X, y - 0.14, CONTENT_W)
        runs = [(lead, {"bold": True, "color": "black"})]
        if rest:
            runs.append((" " + rest, {"color": "body"}))
        return self.text(MARGIN_X, y, 9.9, h, [runs], size=size, anchor="ctr", line=1.4)

    def rule(self, x, y, length, vertical=False):
        return self.rect(x, y, 0.01 if vertical else length, length if vertical else 0.01, "border", None)

    def arrow(self, x, y, w, h, direction="right", color="muted"):
        shape = {"right": MSO_SHAPE.RIGHT_ARROW, "left": MSO_SHAPE.LEFT_ARROW,
                 "down": MSO_SHAPE.DOWN_ARROW, "up": MSO_SHAPE.UP_ARROW}[direction]
        return self.rect(x, y, w, h, color, None, shape=shape)

    def connector(self, x1, y1, x2, y2, color="muted", width_pt=1.0, dashed=False, arrow_end=True, arrow_start=False):
        cn = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        cn.line.color.rgb = rgb(color)
        cn.line.width = Pt(width_pt)
        ln = cn.line._get_or_add_ln()
        if dashed:
            etree.SubElement(ln, qn("a:prstDash")).set("val", "dash")
        if arrow_start:
            etree.SubElement(ln, qn("a:headEnd")).set("type", "triangle")
        if arrow_end:
            etree.SubElement(ln, qn("a:tailEnd")).set("type", "triangle")
        return cn

    # -- diagram helpers (references/diagram-styles.md) --------------------
    def route(self, points, color="muted", width_pt=1.0, dashed=False, arrow_end=True):
        """Orthogonal edge through [(x, y), ...]; only the last segment gets the arrowhead."""
        segs = list(zip(points, points[1:]))
        return [self.connector(a[0], a[1], b[0], b[1], color, width_pt, dashed,
                               arrow_end and i == len(segs) - 1) for i, (a, b) in enumerate(segs)]

    def badge(self, x, y, number, lead="", rest="", size=10, d=0.3):
        """Numbered step header: dark circle with white number, bold caps verb, grey explanation."""
        self.rect(x, y, d, d, "ink", None, shape=MSO_SHAPE.OVAL)
        self.text(x, y, d, d, str(number), size=size, color="white", bold=True, align="ctr", anchor="ctr", wrap=False)
        runs = [(lead.upper(), {"bold": True, "color": "ink"})] if lead else []
        if rest:
            runs.append(((" " if lead else "") + rest, {"color": "body"}))
        if runs:
            self.text(x + d + 0.12, y, 10.0, d, [runs], size=size, anchor="ctr", wrap=False)

    def pills(self, x, y, labels, w=1.1, h=0.4, gap=0.22, variants=None, arrows=True, size=9.5, subs=None):
        """Token/step sequence left to right. variants: one name or a list per pill.
        Returns the (x, y, w, h) box of each pill for brackets, marks and dimensions."""
        n = len(labels)
        widths = w if isinstance(w, (list, tuple)) else [w] * n
        vs = variants if isinstance(variants, (list, tuple)) else [variants or "white"] * n
        boxes, cx = [], x
        for i, lab in enumerate(labels):
            self.node(cx, y, widths[i], h, lab, (subs or [None] * n)[i], vs[i], size=size)
            boxes.append((cx, y, widths[i], h))
            if arrows and i < n - 1:
                self.connector(cx + widths[i] + 0.03, y + h / 2, cx + widths[i] + gap - 0.03, y + h / 2)
            cx += widths[i] + gap
        return boxes

    def bracket(self, x1, x2, y, label="", below=True, color="muted", size=8.5, tick=0.08, label_color=None):
        """Square bracket spanning x1..x2 under (or over) a group, with a centred label."""
        t = tick if below else -tick
        self.connector(x1, y, x1, y + t, color, 0.75, arrow_end=False)
        self.connector(x2, y, x2, y + t, color, 0.75, arrow_end=False)
        self.connector(x1, y + t, x2, y + t, color, 0.75, arrow_end=False)
        if label:
            ly = y + t + 0.03 if below else y + t - 0.25
            self.text(x1, ly, x2 - x1, 0.22, label, size=size, color=label_color or color, bold=True, align="ctr", anchor="ctr")

    def dimension(self, x1, x2, y, label="", color="ink", size=8.5, extend_to=None):
        """Double-headed measure line (latency, window, span) with the label above it.
        extend_to: y of the objects measured; draws thin extension lines up/down to it."""
        if extend_to is not None:
            for xx in (x1, x2):
                self.connector(xx, extend_to, xx, y, "border", 0.75, arrow_end=False)
        self.connector(x1, y, x2, y, color, 1.0, arrow_end=True, arrow_start=True)
        if label:
            self.text(x1, y - 0.26, x2 - x1, 0.22, label, size=size, color=color, bold=True, align="ctr", anchor="ctr")

    def mark(self, x, y, ok=True, size=16):
        """Check (ink) or cross (red) glyph; pair it with a word, colour alone is not the status."""
        return self.text(x, y, 0.3, 0.3, "\u2713" if ok else "\u2717", size=size, color="ink" if ok else "red",
                         bold=True, align="ctr", anchor="ctr", wrap=False)

    def op(self, x, y, symbol="+", d=0.26, size=10):
        """Small operator circle (+, ×, ⊙, σ) placed on a data path."""
        self.rect(x, y, d, d, "white", "muted", 0.75, shape=MSO_SHAPE.OVAL)
        return self.text(x, y, d, d, symbol, size=size, color="ink", bold=True, align="ctr", anchor="ctr", wrap=False)

    def grid(self, x, y, rows, cols, cell=0.14, fill="white", cells=None, stack=0, offset=0.05, border="muted"):
        """Tensor / matrix / vector drawn as a cell grid. cells: {(r, c): fill} highlights
        (use 'pale_red', 'pink', 'red', 'surface'). stack: extra layers drawn behind,
        offset up-right, for the stacked-tensor look. Returns the front grid's box."""
        for k in range(stack, 0, -1):
            self.rect(x + k * offset, y - k * offset, cols * cell, rows * cell, "white", border, 0.5)
        cells = cells or {}
        for r in range(rows):
            for c in range(cols):
                self.rect(x + c * cell, y + r * cell, cell, cell, cells.get((r, c), fill), border, 0.5)
        return (x, y, cols * cell, rows * cell)

    def zoom(self, src, dst, color="muted", outline_src=True):
        """Callout from a small region src=(x, y, w, h) to its enlarged detail panel dst:
        dashed box on the source, thin lines joining the nearer corners."""
        sx, sy, sw, sh = src
        dx, dy, dw, dh = dst
        if outline_src:
            self.rect(sx, sy, sw, sh, None, color, 0.75, dashed=True)
        if dy + dh <= sy:            # detail above
            pairs = [((sx, sy), (dx, dy + dh)), ((sx + sw, sy), (dx + dw, dy + dh))]
        elif dy >= sy + sh:          # detail below
            pairs = [((sx, sy + sh), (dx, dy)), ((sx + sw, sy + sh), (dx + dw, dy))]
        elif dx >= sx + sw:          # detail right
            pairs = [((sx + sw, sy), (dx, dy)), ((sx + sw, sy + sh), (dx, dy + dh))]
        else:                        # detail left
            pairs = [((sx, sy), (dx + dw, dy)), ((sx, sy + sh), (dx + dw, dy + dh))]
        for (ax, ay), (bx, by) in pairs:
            self.connector(ax, ay, bx, by, color, 0.75, arrow_end=False)

    def panel(self, x, y, w, h, label, variant="neutral", size=10):
        """Titled region: a numbered paper section (§3.2 …) or a sub-module group.
        The label sits in a white tab at the top-left. Returns the inner content top y."""
        fill, border, bpt, dashed, tcol, _ = VARIANTS[variant]
        self.rect(x, y, w, h, fill, border, bpt, dashed)
        tw = min(w - 0.3, 0.085 * len(label) + 0.4)
        self.rect(x + 0.12, y + 0.1, tw, 0.26, "white", "border", 0.75)
        self.text(x + 0.12, y + 0.1, tw, 0.26, label, size=size, color="ink" if variant != "dark" else tcol,
                  bold=True, align="ctr", anchor="ctr", wrap=False)
        return y + 0.46

    def span(self, x, y, w, h, fill="pink", label=None, size=8, color="black"):
        """Timeline/Gantt segment (prefill, decode, queue). Pair fills with a legend."""
        self.rect(x, y, w, h, fill, "muted" if fill in ("white", "track") else None, 0.5)
        if label:
            self.text(x, y, w, h, label, size=size, color=color, bold=True, align="ctr", anchor="ctr", wrap=False)

    def legend(self, x, y, items, size=8.5, sw=0.2, gap=0.3):
        """Inline legend: [(variant_or_fill, label), ...] on one row. Returns the end x."""
        cx = x
        for key, label in items:
            if key in VARIANTS:
                fill, border, bpt, dashed, _, _ = VARIANTS[key]
                self.rect(cx, y + 0.02, sw, sw * 0.8, fill, border, bpt or 0.75, dashed)
            else:
                self.rect(cx, y + 0.02, sw, sw * 0.8, key, "muted", 0.5)
            wlab = 0.065 * len(label) + 0.1
            self.text(cx + sw + 0.08, y, wlab, 0.2, label, size=size, color="body", anchor="ctr", wrap=False)
            cx += sw + 0.08 + wlab + gap
        return cx

    def stat(self, x, y, w, value, label, size=26, label_size=10):
        """Big red number with a one-line label under it."""
        return self.text(x, y, w, 0.95, [[(value, {"size": size, "bold": True, "color": "title"})],
                                           [(label, {"size": label_size, "color": "black"})]], line=1.15)

    def bar(self, x, y, w, h, fraction, value_label=None, color="red", track=True, size=9):
        """Horizontal native bar on a light track (for in-slide mini charts)."""
        if track:
            self.rect(x, y, w, h, "track", None)
        bw = max(0.02, w * max(0.0, min(1.0, fraction)))
        self.rect(x, y, bw, h, color, None)
        if value_label:
            self.text(x + bw + 0.08, y, 1.4, h, value_label, size=size, color="black", bold=True, anchor="ctr", wrap=False)

    def table(self, x, y, w, rows, col_widths=None, row_h=0.42, size=10, first_col_bold=True, highlight=None, align=None):
        """Team table: grey header row, hairline row rules, no vertical lines.
        highlight: {(row, col): 'pale'|'accent'}; align: per-column 'l'|'ctr'|'r'."""
        nr, nc = len(rows), len(rows[0])
        gf = self.slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w), Inches(row_h * nr))
        tbl = gf.table
        tblPr = tbl._tbl.tblPr
        for attr in ("firstRow", "bandRow"):
            tblPr.set(attr, "0")
        style = tblPr.find(qn("a:tableStyleId"))
        if style is not None:
            tblPr.remove(style)
        widths = col_widths or [w / nc] * nc
        scale = w / sum(widths)
        for i, cw in enumerate(widths):
            tbl.columns[i].width = Inches(cw * scale)
        highlight = highlight or {}
        for r in range(nr):
            tbl.rows[r].height = Inches(row_h)
            for c in range(nc):
                cell = tbl.cell(r, c)
                cell.margin_left = cell.margin_right = Inches(0.1)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                hl = highlight.get((r, c))
                fill = "surface" if r == 0 else {"pale": "pale_red", "accent": "red"}.get(hl)
                if fill:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = rgb(fill)
                else:
                    cell.fill.background()
                tf = cell.text_frame
                tf.word_wrap = True
                tf.text = ""
                color = "ink" if r == 0 else ("white" if hl == "accent" else ("title" if hl == "pale" else "black"))
                bold = r == 0 or (first_col_bold and c == 0) or hl is not None
                _write(tf, str(rows[r][c]), size, color, bold, (align or ["l"] * nc)[c], line=1.25)
                tcPr = cell._tc.get_or_add_tcPr()
                for side in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                    ln = etree.SubElement(tcPr, qn(side))
                    if side == "a:lnB":
                        ln.set("w", "6350")
                        sf = etree.SubElement(ln, qn("a:solidFill"))
                        etree.SubElement(sf, qn("a:srgbClr")).set("val", C["border"])
                    else:
                        etree.SubElement(ln, qn("a:noFill"))
                # schema order: borders must precede the fill element
                for fill_el in tcPr.findall(qn("a:solidFill")) + tcPr.findall(qn("a:noFill")):
                    if fill_el.getparent() is tcPr:
                        tcPr.remove(fill_el)
                        tcPr.append(fill_el)
        return gf

    def chart(self, x, y, w, h, categories, series, kind="bar", focus=None, number_format="General", legend=True, labels=True, size=8):
        """Native clustered bar/column chart. series: {name: values}; focus series is red, others grey."""
        data = CategoryChartData()
        data.categories = categories
        for name, vals in series.items():
            data.add_series(name, vals)
        ctype = XL_CHART_TYPE.BAR_CLUSTERED if kind == "bar" else XL_CHART_TYPE.COLUMN_CLUSTERED
        gf = self.slide.shapes.add_chart(ctype, Inches(x), Inches(y), Inches(w), Inches(h), data)
        ch = gf.chart
        ch.font.name = FONT
        ch.font.size = Pt(size)
        ch.font.color.rgb = rgb("body")
        greys = ["muted", "grey_bar", "border"]
        gi = 0
        names = list(series)
        for i, s in enumerate(ch.plots[0].series):
            s.format.fill.solid()
            if names[i] == focus or (focus is None and i == len(names) - 1):
                s.format.fill.fore_color.rgb = rgb("red")
            else:
                s.format.fill.fore_color.rgb = rgb(greys[gi % len(greys)])
                gi += 1
        plot = ch.plots[0]
        plot.gap_width = 60
        if labels:
            plot.has_data_labels = True
            plot.data_labels.number_format = number_format
            plot.data_labels.number_format_is_linked = False
            plot.data_labels.font.size = Pt(size)
        ch.has_legend = legend and len(series) > 1
        if ch.has_legend:
            ch.legend.position = XL_LEGEND_POSITION.BOTTOM
            ch.legend.include_in_layout = False
            ch.legend.font.size = Pt(size)
        va = ch.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = rgb("border")
        va.format.line.fill.background()
        va.tick_labels.font.size = Pt(size)
        va.tick_labels.number_format = number_format
        va.tick_labels.number_format_is_linked = False
        ca = ch.category_axis
        ca.format.line.color.rgb = rgb("border")
        ca.tick_labels.font.size = Pt(size)
        return gf

    def image(self, path, x, y, w=None, h=None):
        return self.slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(w) if w else None, Inches(h) if h else None)

    def notes(self, text):
        self.slide.notes_slide.notes_text_frame.text = text

    # -- checks ----------------------------------------------------------
    def lint(self):
        warns = []
        for tb, x, y, w, h, line, wrap in self._text_checks:
            label = tb.text_frame.text.strip().replace("\n", " / ")[:50]
            if x < -0.01 or y < -0.01 or x + w > SLIDE_W + 0.01 or y + h > SLIDE_H + 0.01:
                warns.append(f"outside slide: '{label}'")
            if x + w > LOGO_ZONE[0] and y + h > LOGO_ZONE[1] and label:
                warns.append(f"overlaps footer logo zone: '{label}'")
            need = 0.0
            for p in tb.text_frame.paragraphs:
                sizes = [r.font.size.pt for r in p.runs if r.font.size] or [12]
                if min(sizes) < MIN_PT:
                    warns.append(f"font {min(sizes)} pt below {MIN_PT} pt: '{label}'")
                sz = max(sizes)
                chars = sum(len(r.text) for r in p.runs)
                per_line = max(1, (w * 72) / (sz * 0.55)) if wrap else 1e9
                if not wrap and chars * sz * 0.55 / 72 > w + 0.05:
                    warns.append(f"single-line text likely wider than box ({w:.2f} in): '{label}'")
                need += max(1, math.ceil(chars / per_line)) * sz * line / 72
            if need > h + 0.12:
                warns.append(f"text likely overflows box (needs ~{need:.2f} in, has {h:.2f} in): '{label}'")
        return warns


class TeamDeck:
    def __init__(self, kicker="", corner="triangle", template=TEMPLATE):
        """corner: 'triangle' (plain cyan), a short label such as 'Technology Trends'
        (cyan triangle with rotated white label), or None."""
        self.prs = Presentation(template)
        self.kicker, self.corner = kicker, corner
        self.canvases = []
        self._title_used = False
        self._blank = self._branded_layout("Blank")

    def _branded_layout(self, name):
        """Layout from the same master as the template's title slide (the team's branded master)."""
        master = self.prs.slides[0].slide_layout.slide_master
        for layout in master.slide_layouts:
            if layout.name == name:
                return layout
        raise ValueError(f"layout '{name}' not found under the title slide's master")

    # -- slides ----------------------------------------------------------
    def title_slide(self, title, subtitle="", presenters=(), team=""):
        s = self.prs.slides[0]
        by = {sh.name: sh for sh in s.shapes}

        def fill(shape, paragraphs):
            tf = shape.text_frame
            proto = [copy.deepcopy(p._p) for p in tf.paragraphs]
            for p in list(tf.paragraphs):
                p._p.getparent().remove(p._p)
            txBody = tf._txBody
            for i, text in enumerate(paragraphs):
                p = copy.deepcopy(proto[min(i, len(proto) - 1)])
                runs = p.findall(qn("a:r"))
                for r in runs[1:]:
                    p.remove(r)
                runs[0].find(qn("a:t")).text = text
                txBody.append(p)

        fill(by["Title 3"], [title] + ([subtitle] if subtitle else []))
        fill(by["Text Placeholder 4"], list(presenters) or [" "])
        fill(by["Text Placeholder 5"], [team or " "])
        self._title_used = True
        return s

    def slide(self, title, tagline=None, kicker=None):
        s = self.prs.slides.add_slide(self._blank)
        cv = Canvas(self, s, title)
        self._corner(cv)
        k = self.kicker if kicker is None else kicker
        if k:
            cv.text(1.17, 0.14, 8.0, 0.25, k, size=13, color="kicker", anchor="ctr", wrap=False, line=1.2)
        cv.text(1.17, 0.42, 11.53, 0.47, title, size=26, color="title", bold=True, anchor="ctr", wrap=False, line=1.2)
        if tagline:
            cv.text(1.17, 0.92, 11.53, 0.24, tagline, size=10, color="body", anchor="ctr", wrap=False, line=1.2)
        self.canvases.append(cv)
        return cv

    def contents(self, items, title="Contents"):
        cv = self.slide(title, kicker="")
        nums = [[(f"{i + 1:02d}", {"bold": True, "color": "title"})] for i in range(len(items))]
        cv.text(1.13, 2.0, 0.8, 0.52 * len(items), nums, size=18, line=1.75)
        cv.text(2.08, 2.0, 9.5, 0.52 * len(items), list(items), size=18, color="ink", line=1.75)
        return cv

    def section_divider(self, number, title, summary=""):
        cv = self.slide("", kicker="")
        cv.text(1.17, 2.6, 2.0, 1.0, f"{number:02d}", size=54, color="title", bold=True, anchor="ctr", wrap=False)
        cv.text(3.0, 2.6, 9.0, 1.0, title, size=30, color="ink", bold=True, anchor="ctr")
        if summary:
            cv.text(3.0, 3.6, 8.5, 0.9, summary, size=14, color="body")
        return cv

    def _corner(self, cv):
        if not self.corner:
            return
        if self.corner == "triangle":
            tri = cv.rect(0, 0, 1.06, 0.94, "corner", None, shape=MSO_SHAPE.RIGHT_TRIANGLE)
            tri._element.spPr.find(qn("a:xfrm")).set("flipV", "1")
        else:
            tri = cv.rect(0, 0, 0.79, 0.86, "corner", "corner", 0.75, shape=MSO_SHAPE.RIGHT_TRIANGLE)
            tri._element.spPr.find(qn("a:xfrm")).set("flipV", "1")
            lab = self.prs.slides[-1].shapes.add_textbox(Inches(-0.12), Inches(0.10), Inches(0.92), Inches(0.34))
            _frame(lab.text_frame, "ctr", wrap=True)
            _write(lab.text_frame, self.corner, 7, "white", True, "ctr", 1.0)
            for r in lab.text_frame.paragraphs[0].runs:
                r.font.name = "Arial"
            lab.rotation = -45

    # -- output ----------------------------------------------------------
    def lint(self):
        warns = []
        for i, cv in enumerate(self.canvases, start=2 if self._title_used else 1):
            if len(cv.name) > 52:
                warns.append(f"slide {i}: title has {len(cv.name)} chars; titles over ~52 chars run past the slide at 26 pt")
            warns += [f"slide {i}: {w}" for w in cv.lint()]
        if not self._title_used:
            warns.append("title slide still has template placeholder text; call title_slide() or delete slide 1")
        return warns

    def save(self, path, quiet=False):
        os.makedirs(os.path.dirname(os.path.abspath(path)), exist_ok=True)
        self.prs.save(path)
        warns = self.lint()
        if not quiet:
            print(f"saved {path} ({len(self.prs.slides)} slides)")
            for w in warns:
                print("WARN", w, file=sys.stderr)
        return warns
